# scripts/tool/ghi_excel.py
# -*- coding: utf-8 -*-
"""Ghi file kết quả: sheet 'Upload' (27 cột)."""
import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Đúng thứ tự template. Cột trống giữa "Số lượng" và "Storage location" (khoá
# nội bộ "H") đã bỏ khỏi file output theo yêu cầu khách 25/09 -> không xuất hiện
# trong TEN_COT/KHOA, dù dung_bang.py vẫn còn giữ khoá "H" nội bộ (luôn rỗng).
# "Ghi chú đơn hàng" (khoá nội bộ "Z") nằm ngoài template gốc, thêm 25/09 theo
# yêu cầu khách — hiện CHƯA có nguồn data, luôn để trống (dung_bang.py không
# gán giá trị cho khoá "Z").
TEN_COT = [
    "Mã ODO", "Mã đơn hàng", "Mã đặt hàng", "Ngày gửi y/c", "Mã hàng hóa",
    "Tên hàng hóa/đơn vị tính", "Số lượng", "Storage location",
    "Mã điểm lấy hàng", "Mã Sold to", "Mã Ship to", "Tên Ship to", "Địa chỉ",
    "Tỉnh điểm giao hàng", "Quận điểm giao hàng", "Mã hình thức thanh toán",
    "Giá chuyến COD", "Khách hàng trả phí vận chuyển (Có/ Không)", "ETD", "ETA",
    "Mã HU", "Ngày bán hàng", "Ghi chú đơn hàng", "Loại vận chuyển", "Loại dịch vụ",
    "Thông tin người nhận", "Lý do cảnh báo",
]

# Z ("Ghi chú đơn hàng") chèn giữa W và X -> không dùng range chữ cái liên tục
# được nữa, phải liệt kê tay đúng thứ tự cột thật. "AB" (cuối, "Lý do cảnh báo")
# không phải khoá thường có trong dong[i] -> ghi riêng từ r["_ly_do"], xem _sheet_upload.
KHOA = [k for k in "ABCDEFGHIJKLMNOPQRSTUVW" if k != "H"] + ["Z", "X", "Y", "AA", "AB"]

# Sheet ẩn mang Mã Sold to -> Tên Sold to (DO."Sold-to party") sang Bước 2
# (stm_soldto_shipto.xu_ly/stm_api.ghi_file_soldto_shipto) - file upload đơn hàng KHÔNG có
# cột nào chứa tên Sold-to (template chỉ có mã), nên không thể thêm vào sheet "Upload" mà
# STM đọc. Khách chốt 20/10/2026: dùng sheet ẩn trong CHÍNH file này thay vì file .json rời
# (dễ bị khách làm thất lạc) hoặc thêm cột hiển thị (phải hỏi lại STM có chấp nhận không).
SHEET_TEN_SOLDTO = "_TenSoldTo"
COT_TEN_SOLDTO = ["Mã Sold to", "Tên Sold to"]
VANG = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
DINH_DANG_NGAY = {"D": "dd/mm/yyyy", "T": "dd/mm/yyyy hh:mm",
                  "U": "dd/mm/yyyy hh:mm", "W": "dd/mm/yyyy hh:mm"}
RONG_COT = {"M": 38, "N": 60, "O": 20, "P": 20, "D": 13, "T": 17, "U": 17, "W": 17, "AA": 45}

# Khách không rành flow tool, chỉ cần biết: cái nào TỰ SỬA được trên Excel,
# cái nào phải BÁO LẠI vì thiếu ngay từ dữ liệu SAP gốc (tool không có gì để suy).
NHOM_SAP = "Thiếu dữ liệu SAP gốc"
NHOM_DIA_CHI = "Cần soát địa chỉ"
NHOM_STM = "Sold-to/Ship-to trên STM"
NHOM_KHAC = "Khác"
GIAI_THICH_NHOM = {
    NHOM_SAP: "không tự sửa trên Excel được — báo lại bộ phận SAP",
    NHOM_DIA_CHI: "có thể sửa trực tiếp trên Excel",
    NHOM_STM: "tool chưa tạo được trên STM — khách tự tạo tay hoặc chạy lại",
    NHOM_KHAC: "xem cột Chi tiết",
}
_LOAI_TO_NHOM = {
    "Thiếu ngày gửi / ETD / ETA (cột D, T, U)": NHOM_SAP,
    "Số lượng ≤ 0": NHOM_SAP,
    "Tỉnh SAP và tỉnh theo toạ độ khác nhau": NHOM_DIA_CHI,
    "Địa chỉ cần soát": NHOM_DIA_CHI,
    "Không xác định được Quận/Huyện": NHOM_DIA_CHI,
    "Quận không có trong danh sách STM chấp nhận": NHOM_DIA_CHI,
    "Không xác định được Tỉnh": NHOM_DIA_CHI,
    "Tỉnh độ tin cậy thấp": NHOM_DIA_CHI,
    "Không tạo được Sold-to/Ship-to mới trên STM": NHOM_STM,
}


def nhom_cua(loai):
    return _LOAI_TO_NHOM.get(loai, NHOM_KHAC)


def don_theo_nhom(canh_bao):
    """-> {nhóm: set(mã ODO)}, sắp xếp nhóm đông đơn nhất trước."""
    theo_nhom = {}
    for c in canh_bao:
        theo_nhom.setdefault(nhom_cua(c["loai"]), set()).add(c["ma_odo"])
    return dict(sorted(theo_nhom.items(), key=lambda kv: -len(kv[1])))


def ghi_file(duong_dan, dong):
    wb = openpyxl.Workbook()
    _sheet_upload(wb.active, dong)
    _sheet_ten_soldto(wb, dong)
    try:
        wb.save(duong_dan)
    except PermissionError:
        raise PermissionError(
            f"Không ghi được '{duong_dan}'. File đang mở trong Excel — hãy đóng lại rồi chạy lại."
        )
    finally:
        wb.close()


def _sheet_upload(ws, dong):
    ws.title = "Upload"
    ws.append(TEN_COT)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for r in dong:
        gia_tri = [r.get(k, "") for k in KHOA[:-1]]
        chi_tiet_theo_loai = r.get("_chi_tiet", {})
        ly_do_day_du = [
            f"{loai} — {chi_tiet_theo_loai[loai]}" if loai in chi_tiet_theo_loai else loai
            for loai in r.get("_ly_do", [])
        ]
        gia_tri.append("; ".join(ly_do_day_du))
        ws.append(gia_tri)
        i = ws.max_row
        thieu = set(r.get("_thieu", []))
        for j, k in enumerate(KHOA, start=1):
            o = ws.cell(i, j)
            if k in thieu:
                o.fill = VANG
            if k in DINH_DANG_NGAY and isinstance(o.value, datetime.datetime):
                o.number_format = DINH_DANG_NGAY[k]

    for k, w in RONG_COT.items():
        ws.column_dimensions[get_column_letter(KHOA.index(k) + 1)].width = w


def _sheet_ten_soldto(wb, dong):
    """Ghi sheet ẩn (xem SHEET_TEN_SOLDTO) - mỗi Mã Sold to 1 dòng, lấy tên THẬT đầu tiên
    gặp (r['_ten_soldto'], xem dung_bang.py). Thiếu khoá/tên rỗng -> vẫn ghi dòng với tên
    rỗng, không crash (Bước 2 khi đó để trống, giống hành vi cũ)."""
    ten_theo_ma = {}
    for r in dong:
        ma = r.get("K", "")
        if not ma or ma in ten_theo_ma:
            continue
        ten_theo_ma[ma] = r.get("_ten_soldto", "")
    ws = wb.create_sheet(SHEET_TEN_SOLDTO)
    ws.append(COT_TEN_SOLDTO)
    for ma, ten in ten_theo_ma.items():
        ws.append([ma, ten])
    ws.sheet_state = "hidden"


def _doc_ten_soldto(wb):
    """-> {ma_soldto: ten}. Không có sheet (file cũ trước khi có tính năng này) -> {}."""
    if SHEET_TEN_SOLDTO not in wb.sheetnames:
        return {}
    ws = wb[SHEET_TEN_SOLDTO]
    ket_qua = {}
    for i in range(2, ws.max_row + 1):
        ma = ws.cell(i, 1).value
        if ma:
            ket_qua[ma] = ws.cell(i, 2).value or ""
    return ket_qua


def doc_lai_upload(duong_dan):
    """Đọc lại sheet 'Upload' đã xuất (có thể đã được khách sửa tay trong Excel) -> list[dict]
    theo khoá K/L/M/N/O/P + '_ten_soldto' (tra theo K từ sheet ẩn SHEET_TEN_SOLDTO, xem
    _sheet_ten_soldto), kèm '_hang' (số dòng Excel thật, bắt đầu từ 2) để `ghi_de_upload()`
    biết ghi đè đúng ô. Dùng cho Bước 2 (main.tao_file_soldto_shipto) - 'M' (Tên Ship to) và
    '_ten_soldto' hiện không còn nơi nào tiêu thụ sau khi bỏ file upload_soldto_shipto riêng
    (21/10/2026, xem CLAUDE.md) - giữ lại để không vỡ hợp đồng dữ liệu, có thể dọn ở phiên sau
    nếu xác nhận chắc chắn không còn dùng."""
    wb = openpyxl.load_workbook(duong_dan, data_only=True)
    try:
        ws = wb["Upload"]
        ten_soldto = _doc_ten_soldto(wb)
        ket_qua = []
        so_cot = len(KHOA)
        for i in range(2, ws.max_row + 1):
            if all(ws.cell(i, j).value in (None, "") for j in range(1, so_cot + 1)):
                continue
            d = {"_hang": i}
            for k in ("K", "L", "M", "N", "O", "P"):
                v = ws.cell(i, KHOA.index(k) + 1).value
                d[k] = "" if v is None else v
            d["_ten_soldto"] = ten_soldto.get(d["K"], "")
            ket_qua.append(d)
        return ket_qua
    finally:
        wb.close()


def ghi_de_upload(duong_dan, thay_doi):
    """Ghi đè lại đúng ô L/N/O/P của file `duong_dan` theo `thay_doi` (list dict có '_hang' +
    bất kỳ khoá nào trong 'L','N','O','P' cần sửa) - không đụng ô nào khác, không đụng cột
    'Lý do cảnh báo'. Dùng ở Bước 2, sau khi `stm_soldto_shipto.xu_ly()` xử lý xong."""
    if not thay_doi:
        return
    wb = openpyxl.load_workbook(duong_dan)
    try:
        ws = wb["Upload"]
        for d in thay_doi:
            for k in ("L", "N", "O", "P"):
                if k in d:
                    ws.cell(d["_hang"], KHOA.index(k) + 1).value = d[k]
        try:
            wb.save(duong_dan)
        except PermissionError:
            raise PermissionError(
                f"Không ghi được '{duong_dan}'. File đang mở trong Excel — hãy đóng lại rồi "
                "bấm lại nút 3."
            )
    finally:
        wb.close()


def xoa_canh_bao_dia_chi(duong_dan, cac_hang):
    """Dùng ở Bước 2 (main.tao_file_soldto_shipto), sau khi quét lại Tỉnh/Quận (trực tiếp
    hợp lệ hoặc tự sửa/suy lại qua tra_dia_chi) thấy dòng đã đúng: xoá các mục cảnh báo
    thuộc NHOM_DIA_CHI khỏi cột 'Lý do cảnh báo' + bỏ tô vàng cột Địa chỉ/Tỉnh/Quận (N/O/P)
    của đúng các dòng Excel trong `cac_hang` (số dòng thật, xem doc_lai_upload). Không đụng
    cảnh báo nhóm khác (vd 'Thiếu Mã đơn hàng') hay dòng không nằm trong `cac_hang`."""
    if not cac_hang:
        return
    wb = openpyxl.load_workbook(duong_dan)
    try:
        ws = wb["Upload"]
        col_ab = KHOA.index("AB") + 1
        for hang in cac_hang:
            o = ws.cell(hang, col_ab)
            if o.value:
                con_lai = [p for p in str(o.value).split("; ")
                           if p and nhom_cua(p.split(" — ")[0]) != NHOM_DIA_CHI]
                o.value = "; ".join(con_lai)
            for k in ("N", "O", "P"):
                ws.cell(hang, KHOA.index(k) + 1).fill = PatternFill(fill_type=None)
        try:
            wb.save(duong_dan)
        except PermissionError:
            raise PermissionError(
                f"Không ghi được '{duong_dan}'. File đang mở trong Excel — hãy đóng lại rồi "
                "bấm lại nút 3."
            )
    finally:
        wb.close()
